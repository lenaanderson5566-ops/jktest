from __future__ import annotations

from io import BytesIO

from django.contrib import admin, messages
from django.http import HttpRequest, HttpResponse
from django.middleware.csrf import get_token
from django.shortcuts import redirect, render
from django.urls import path, reverse
from openpyxl import Workbook, load_workbook

from .models import (
    GlobalConfig,
    Organization,
    DenominationPackagingSpec,
    PackingStation,
    StationDenominationEfficiency,
    StationDenominationSupport,
    TransferSegment,
    TransportRoute,
)


class ExcelMixin:
    model_label = ''

    def upload_page(self, request: HttpRequest):
        if request.method == 'POST' and request.FILES.get('file'):
            try:
                self.handle_upload(request.FILES['file'])
                messages.success(request, f'{self.model_label} 导入成功')
            except Exception as exc:  # noqa: BLE001
                messages.error(request, f'{self.model_label} 导入失败: {exc}')
            return redirect('..')

        token = get_token(request)
        return HttpResponse(
            '<h3>Excel 导入</h3>'
            '<form method="post" enctype="multipart/form-data">'
            f'<input type="hidden" name="csrfmiddlewaretoken" value="{token}" />'
            f'<p>{self.model_label} 导入文件：</p>'
            '<input type="file" name="file" accept=".xlsx" required />'
            '<button type="submit">上传并导入</button>'
            '</form>'
        )


@admin.register(TransportRoute)
class TransportRouteAdmin(admin.ModelAdmin, ExcelMixin):
    change_list_template = 'admin/excel_change_list.html'
    list_display = ('route_no', 'route_name', 'enabled')
    search_fields = ('route_no', 'route_name')
    model_label = '押运线路'

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('import-excel/', self.admin_site.admin_view(self.upload_page), name='masterdata_transportroute_import'),
            path('export-excel/', self.admin_site.admin_view(self.export_excel), name='masterdata_transportroute_export'),
            path('template-excel/', self.admin_site.admin_view(self.template_excel), name='masterdata_transportroute_template'),
        ]
        return custom_urls + urls

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context.update({
            'excel_import_url': reverse('admin:masterdata_transportroute_import'),
            'excel_export_url': reverse('admin:masterdata_transportroute_export'),
            'excel_template_url': reverse('admin:masterdata_transportroute_template'),
        })
        return super().changelist_view(request, extra_context)

    @staticmethod
    def handle_upload(file_obj):
        wb = load_workbook(file_obj)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            TransportRoute.objects.update_or_create(
                route_no=str(row[0]).strip(),
                defaults={
                    'route_name': str(row[1] or '').strip(),
                    'enabled': bool(row[2]) if row[2] is not None else True,
                    'remark': str(row[3] or '').strip(),
                },
            )

    def export_excel(self, request):
        wb = Workbook()
        ws = wb.active
        ws.append(['线路号', '线路名称', '是否启用', '备注'])
        for obj in TransportRoute.objects.all().order_by('route_no'):
            ws.append([obj.route_no, obj.route_name, obj.enabled, obj.remark])
        return self._wb_response(wb, 'transport_route_export.xlsx')

    def template_excel(self, request):
        wb = Workbook()
        ws = wb.active
        ws.append(['线路号', '线路名称', '是否启用', '备注'])
        ws.append(['R001', '示例线路', True, '样例数据'])
        return self._wb_response(wb, 'transport_route_template.xlsx')

    @staticmethod
    def _wb_response(wb, filename):
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename={filename}'
        return response


@admin.register(Organization)
class OrganizationAdmin(admin.ModelAdmin, ExcelMixin):
    change_list_template = 'admin/excel_change_list.html'
    list_display = ('org_no', 'org_name', 'route', 'enabled')
    search_fields = ('org_no', 'org_name')
    list_filter = ('route', 'enabled')
    model_label = '机构'

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('import-excel/', self.admin_site.admin_view(self.upload_page), name='masterdata_organization_import'),
            path('export-excel/', self.admin_site.admin_view(self.export_excel), name='masterdata_organization_export'),
            path('template-excel/', self.admin_site.admin_view(self.template_excel), name='masterdata_organization_template'),
        ]
        return custom_urls + urls

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context.update({
            'excel_import_url': reverse('admin:masterdata_organization_import'),
            'excel_export_url': reverse('admin:masterdata_organization_export'),
            'excel_template_url': reverse('admin:masterdata_organization_template'),
        })
        return super().changelist_view(request, extra_context)

    @staticmethod
    def handle_upload(file_obj):
        wb = load_workbook(file_obj)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            route_no = str(row[2] or '').strip()
            route = TransportRoute.objects.filter(route_no=route_no).first()
            if not route:
                raise ValueError(f'找不到线路号: {route_no}')
            Organization.objects.update_or_create(
                org_no=str(row[0]).strip(),
                defaults={
                    'org_name': str(row[1] or '').strip(),
                    'route': route,
                    'enabled': bool(row[3]) if row[3] is not None else True,
                    'remark': str(row[4] or '').strip(),
                },
            )

    def export_excel(self, request):
        wb = Workbook()
        ws = wb.active
        ws.append(['机构号', '机构名称', '所属线路号', '是否启用', '备注'])
        for obj in Organization.objects.select_related('route').all().order_by('org_no'):
            ws.append([obj.org_no, obj.org_name, obj.route.route_no, obj.enabled, obj.remark])
        return TransportRouteAdmin._wb_response(wb, 'organization_export.xlsx')

    def template_excel(self, request):
        wb = Workbook()
        ws = wb.active
        ws.append(['机构号', '机构名称', '所属线路号', '是否启用', '备注'])
        ws.append(['ORG001', '示例机构', 'R001', True, '样例数据'])
        return TransportRouteAdmin._wb_response(wb, 'organization_template.xlsx')


@admin.register(DenominationPackagingSpec)
class DenominationPackagingSpecAdmin(admin.ModelAdmin):
    list_display = ('currency_type', 'denomination', 'units_per_package', 'enabled')
    list_filter = ('currency_type', 'enabled')


@admin.register(PackingStation)
class PackingStationAdmin(admin.ModelAdmin):
    list_display = ('station_no', 'station_name', 'station_order', 'station_type', 'enabled')
    list_filter = ('station_type', 'enabled')


@admin.register(GlobalConfig)
class GlobalConfigAdmin(admin.ModelAdmin):
    list_display = ('config_key', 'config_value', 'enabled', 'remark')
    list_filter = ('enabled',)
    search_fields = ('config_key', 'remark')


def build_pipeline_context():
    stations = list(PackingStation.objects.filter(enabled=True).order_by('station_order'))
    segments = list(TransferSegment.objects.filter(enabled=True).select_related('from_station', 'to_station'))
    seg_map = {(s.from_station_id, s.to_station_id): s for s in segments}

    support_rows = StationDenominationSupport.objects.filter(enabled=True).order_by('denomination')
    support_map = {}
    for row in support_rows:
        support_map.setdefault(row.station_id, []).append(f"{row.get_currency_type_display()} {row.denomination}")

    eff_rows = StationDenominationEfficiency.objects.filter(enabled=True)
    eff_map = {}
    for row in eff_rows:
        eff_map.setdefault(row.station_id, {})[str(row.denomination)] = float(row.unit_boxing_seconds)

    chains = []
    for idx, station in enumerate(stations):
        next_station = stations[idx + 1] if idx + 1 < len(stations) else None
        supports = support_map.get(station.id, [])
        effective_unit_times = []
        for text in supports:
            denom = text.split()[-1]
            unit_time = eff_map.get(station.id, {}).get(denom, float(station.unit_boxing_seconds))
            effective_unit_times.append(f"{text}: {unit_time}s")
        if not effective_unit_times:
            effective_unit_times.append(f"默认: {float(station.unit_boxing_seconds)}s")

        chains.append({
            'station': station,
            'segment': seg_map.get((station.id, next_station.id)) if next_station else None,
            'supports': supports,
            'effective_unit_times': effective_unit_times,
            'effective_fixed_time': float(station.fixed_boxing_seconds),
        })
    return {'chains': chains, 'segments': segments}


def pipeline_overview_page(request, admin_site):
    from apps.orders.optimization import build_strategy_comparison

    strategy_context = build_strategy_comparison()
    return render(request, 'admin/pipeline_overview.html', {
        **admin_site.each_context(request),
        'title': '流水线概览示意图',
        **build_pipeline_context(),
        **strategy_context,
        'strategy_export_url': reverse('admin_strategy_export'),
    })


def pipeline_overview_ortools_page(request, admin_site):
    from apps.orders.optimization_ortools import build_strategy_comparison_ortools

    started = request.GET.get('run') == '1'
    strategy_context = (
        build_strategy_comparison_ortools()
        if started else
        {
            'has_data': False,
            'results': [],
            'summary': {'total_organizations': 0, 'total_boxes': 0, 'denomination_bundles': []},
        }
    )
    return render(request, 'admin/pipeline_overview.html', {
        **admin_site.each_context(request),
        'title': '流水线概览示意图（ORTools）',
        **build_pipeline_context(),
        **strategy_context,
        'strategy_export_url': reverse('admin_strategy_export_ortools'),
        'ortools_lazy': True,
        'ortools_started': started,
    })


def _build_strategy_export_workbook(results, target):
    if target:
        results = [row for row in results if row.key == target]

    wb = Workbook()
    ws = wb.active
    ws.title = '方案汇总'
    ws.append([
        '方案key', '方案名称', '总完工时长(秒)', '线路连续性', '线路切换次数', '线路聚集度', '排序预览'
    ])

    for row in results:
        metrics = row.metrics
        ws.append([
            row.key,
            row.name,
            metrics.get('total_finish_seconds', 0),
            metrics.get('line_continuity', 0),
            metrics.get('route_switches', 0),
            metrics.get('route_cluster_score', 0),
            ' -> '.join(metrics.get('sequence_preview', [])),
        ])

    for row in results:
        metrics = row.metrics
        detail_rows = metrics.get('box_details', [])

        ws_box = wb.create_sheet(title=f"{row.key[:20]}_箱序")
        ws_box.append([
            '箱序', '订单号', '机构', '线路', '箱号', '投放时间', '完成时间', '总捆数', '面额明细'
        ])
        for detail in detail_rows:
            denom_txt = '; '.join([f"{k}:{v}" for k, v in sorted(detail.get('denoms', {}).items())])
            ws_box.append([
                detail.get('box_index'),
                detail.get('order_no'),
                detail.get('organization_name'),
                detail.get('route_no'),
                detail.get('box_seq_no'),
                detail.get('release_time'),
                detail.get('finish_time'),
                detail.get('total_bundles'),
                denom_txt,
            ])

        ws_station = wb.create_sheet(title=f"{row.key[:20]}_工位")
        ws_station.append([
            '箱序', '订单号', '机构', '线路', '工位', '开始时间', '结束时间', '处理时长', '分配明细'
        ])
        for detail in detail_rows:
            for step in detail.get('station_steps', []):
                alloc_txt = '; '.join([f"{k}:{v}" for k, v in sorted(step.get('allocations', {}).items())])
                ws_station.append([
                    detail.get('box_index'),
                    detail.get('order_no'),
                    detail.get('organization_name'),
                    detail.get('route_no'),
                    step.get('station_name'),
                    step.get('start'),
                    step.get('end'),
                    step.get('process_seconds'),
                    alloc_txt,
                ])
    return wb


def pipeline_strategy_export_page(request, admin_site):
    from apps.orders.optimization import build_strategy_comparison

    target = (request.GET.get('strategy') or '').strip()
    context = build_strategy_comparison(include_details=True)
    results = context.get('results', [])
    wb = _build_strategy_export_workbook(results, target)
    filename = f"strategy_detail_{target or 'all'}.xlsx"
    return TransportRouteAdmin._wb_response(wb, filename)


def pipeline_strategy_export_ortools_page(request, admin_site):
    from apps.orders.optimization_ortools import build_strategy_comparison_ortools

    target = (request.GET.get('strategy') or '').strip()
    context = build_strategy_comparison_ortools(include_details=True)
    results = context.get('results', [])
    wb = _build_strategy_export_workbook(results, target)
    filename = f"strategy_detail_{target or 'all'}.xlsx"
    return TransportRouteAdmin._wb_response(wb, filename)


@admin.register(StationDenominationSupport)
class StationDenominationSupportAdmin(admin.ModelAdmin):
    list_display = ('station_name', 'get_currency_type_display', 'denomination', 'enabled')
    list_filter = ('currency_type', 'enabled')
    search_fields = ('station__station_name', 'station__station_no')

    @admin.display(description='工位名称')
    def station_name(self, obj):
        return obj.station.station_name


admin.site.register(StationDenominationEfficiency)
admin.site.register(TransferSegment)
