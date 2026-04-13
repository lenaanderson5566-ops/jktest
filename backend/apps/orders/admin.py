from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.contrib import admin, messages
from django.http import HttpResponse
from django.middleware.csrf import get_token
from django.shortcuts import redirect
from django.urls import path, reverse
from openpyxl import Workbook, load_workbook

from apps.masterdata.models import (
    CurrencyType,
    DenominationPackagingSpec,
    Organization,
    StationDenominationSupport,
    TransportRoute,
)
from .models import OrganizationOrder


DENOMINATION_TO_FIELD = {
    '100': 'qty_100',
    '50': 'qty_50',
    '20': 'qty_20',
    '10': 'qty_10',
    '5': 'qty_5',
    '1': 'qty_coin_1',
    '0.5': 'qty_coin_05',
    '0.1': 'qty_coin_01',
}

FIELD_TO_DENOMINATION = {v: k for k, v in DENOMINATION_TO_FIELD.items()}
DENOMINATION_LABEL_MAP = {
    '100': '100元捆数',
    '50': '50元捆数',
    '20': '20元捆数',
    '10': '10元捆数',
    '5': '5元捆数',
    '1': '1元包数',
    '0.5': '0.5元包数',
    '0.1': '0.1元包数',
}
DEFAULT_ADMIN_FIELDS = ('order_no', 'order_date', 'organization', 'route')
TAIL_ADMIN_FIELDS = ('status', 'remark')


@admin.register(OrganizationOrder)
class OrganizationOrderAdmin(admin.ModelAdmin):
    change_list_template = 'admin/excel_change_list.html'
    list_display = ('order_no', 'order_date', 'organization', 'route', 'status')
    search_fields = ('order_no', 'organization__org_no', 'organization__org_name')
    list_filter = ('order_date', 'route', 'status')

    def get_fields(self, request, obj=None):
        configured_denominations = _get_configured_denominations()
        qty_fields = [DENOMINATION_TO_FIELD[denom] for denom in configured_denominations]
        if not qty_fields:
            qty_fields = list(DENOMINATION_TO_FIELD.values())
        return [*DEFAULT_ADMIN_FIELDS, *qty_fields, *TAIL_ADMIN_FIELDS]

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('import-excel/', self.admin_site.admin_view(self.import_excel), name='orders_order_import'),
            path('export-excel/', self.admin_site.admin_view(self.export_excel), name='orders_order_export'),
            path('template-excel/', self.admin_site.admin_view(self.template_excel), name='orders_order_template'),
        ]
        return custom_urls + urls

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context.update({
            'excel_import_url': reverse('admin:orders_order_import'),
            'excel_export_url': reverse('admin:orders_order_export'),
            'excel_template_url': reverse('admin:orders_order_template'),
        })
        return super().changelist_view(request, extra_context)

    def import_excel(self, request):
        if request.method == 'POST' and request.FILES.get('file'):
            try:
                self._handle_upload(request.FILES['file'])
                messages.success(request, '机构订单导入成功（自动兼容长表/宽表）')
            except Exception as exc:  # noqa: BLE001
                messages.error(request, f'机构订单导入失败: {exc}')
            return redirect('..')

        token = get_token(request)
        return HttpResponse(
            '<h3>机构订单 Excel 导入（兼容长表/宽表，长表支持金额自动换算包/捆数）</h3>'
            '<form method="post" enctype="multipart/form-data">'
            f'<input type="hidden" name="csrfmiddlewaretoken" value="{token}" />'
            '<input type="file" name="file" accept=".xlsx" required />'
            '<button type="submit">上传并导入</button>'
            '</form>'
        )

    @classmethod
    def _handle_upload(cls, file_obj):
        wb = load_workbook(file_obj)
        ws = wb.active
        headers = [str(v).strip() if v is not None else '' for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

        wide_required = {'订单编号', '订单日期', '机构号', '线路号', '100元捆数', '0.1元包数'}
        long_required = {'订单日期', '机构号', '线路号', '面额'}
        header_set = set(headers)

        if wide_required.issubset(header_set):
            cls._import_wide(ws, headers)
            return
        if long_required.issubset(header_set):
            cls._import_long(ws, headers)
            return
        raise ValueError('无法识别模板格式，请使用系统下载的宽表或长表样表。')

    @staticmethod
    def _import_wide(ws, headers):
        idx = {name: i for i, name in enumerate(headers)}
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                if not row or not row[idx['订单编号']]:
                    continue
                order_no = str(row[idx['订单编号']]).strip()
                order_date = _to_date(row[idx['订单日期']])
                org_no = str(row[idx['机构号']] or '').strip()
                route_no = str(row[idx['线路号']] or '').strip()
                org = Organization.objects.filter(org_no=org_no).first()
                route = TransportRoute.objects.filter(route_no=route_no).first()
                if not org:
                    raise ValueError(f'找不到机构号: {org_no}，请先在机构主数据维护')
                if not route:
                    raise ValueError(f'找不到线路号: {route_no}，请先在押运线路主数据维护')

                payload = {
                'order_date': order_date,
                'organization': org,
                'route': route,
                'qty_100': int(row[idx.get('100元捆数', -1)] or 0),
                'qty_50': int(row[idx.get('50元捆数', -1)] or 0),
                'qty_20': int(row[idx.get('20元捆数', -1)] or 0),
                'qty_10': int(row[idx.get('10元捆数', -1)] or 0),
                'qty_5': int(row[idx.get('5元捆数', -1)] or 0),
                'qty_coin_1': int(row[idx.get('1元包数', -1)] or 0),
                'qty_coin_05': int(row[idx.get('0.5元包数', -1)] or 0),
                'qty_coin_01': int(row[idx.get('0.1元包数', -1)] or 0),
                'status': str(row[idx.get('订单状态', -1)] or 'NEW'),
                'remark': str(row[idx.get('备注', -1)] or '').strip(),
            }
                for field, qty in payload.items():
                    if field in FIELD_TO_DENOMINATION and qty:
                        _validate_denomination_bound(FIELD_TO_DENOMINATION[field])

                OrganizationOrder.objects.update_or_create(order_no=order_no, defaults=payload)
            except Exception as exc:  # noqa: BLE001
                raise ValueError(f'宽表第{row_no}行导入失败：{exc}') from exc

    @staticmethod
    def _import_long(ws, headers):
        idx = {name: i for i, name in enumerate(headers)}
        grouped = defaultdict(lambda: {
            'qty_100': 0,
            'qty_50': 0,
            'qty_20': 0,
            'qty_10': 0,
            'qty_5': 0,
            'qty_coin_1': 0,
            'qty_coin_05': 0,
            'qty_coin_01': 0,
            'status': 'NEW',
            'remark': '',
        })

        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                if not row:
                    continue
                order_date = _to_date(row[idx['订单日期']])
                org_no = str(row[idx['机构号']] or '').strip()
                route_no = str(row[idx['线路号']] or '').strip()
                if not org_no or not route_no:
                    continue

                order_no_col = idx.get('订单编号')
                order_no_val = str(row[order_no_col]).strip() if order_no_col is not None and row[order_no_col] else ''
                order_no = order_no_val or f'ORD-{order_date.strftime("%Y%m%d")}-{org_no}-{route_no}'
                key = (order_no, order_date, org_no, route_no)

                denom = _normalize_denomination(row[idx['面额']])
                field = DENOMINATION_TO_FIELD.get(denom)
                if not field:
                    raise ValueError(f'不支持的面额: {row[idx["面额"]]}')
                _validate_denomination_bound(denom)
                qty_col = idx.get('数量')
                amount_col = idx.get('金额')
                qty_raw = row[qty_col] if qty_col is not None else None
                amount_raw = row[amount_col] if amount_col is not None else None
                qty = _resolve_quantity_from_row(denom, qty_raw, amount_raw)
                grouped[key][field] += qty

                status_col = idx.get('订单状态')
                if status_col is not None and row[status_col]:
                    grouped[key]['status'] = str(row[status_col]).strip()
                remark_col = idx.get('备注')
                if remark_col is not None and row[remark_col]:
                    grouped[key]['remark'] = str(row[remark_col]).strip()
            except Exception as exc:  # noqa: BLE001
                raise ValueError(f'长表第{row_no}行导入失败：{exc}') from exc

        for (order_no, order_date, org_no, route_no), payload in grouped.items():
            org = Organization.objects.filter(org_no=org_no).first()
            route = TransportRoute.objects.filter(route_no=route_no).first()
            if not org:
                raise ValueError(f'聚合后订单 {order_no} 找不到机构号: {org_no}，请先维护机构主数据')
            if not route:
                raise ValueError(f'聚合后订单 {order_no} 找不到线路号: {route_no}，请先维护押运线路主数据')
            OrganizationOrder.objects.update_or_create(
                order_no=order_no,
                defaults={
                    'order_date': order_date,
                    'organization': org,
                    'route': route,
                    **payload,
                },
            )

    def export_excel(self, request):
        wb = Workbook()
        ws = wb.active
        ws.title = '宽表导出'
        ws.append([
            '订单编号', '订单日期', '机构号', '线路号',
            '100元捆数', '50元捆数', '20元捆数', '10元捆数', '5元捆数',
            '1元包数', '0.5元包数', '0.1元包数', '订单状态', '备注',
        ])
        for obj in OrganizationOrder.objects.select_related('organization', 'route').all().order_by('order_date', 'order_no'):
            ws.append([
                obj.order_no, obj.order_date, obj.organization.org_no, obj.route.route_no,
                obj.qty_100, obj.qty_50, obj.qty_20, obj.qty_10, obj.qty_5,
                obj.qty_coin_1, obj.qty_coin_05, obj.qty_coin_01, obj.status, obj.remark,
            ])
        return self._wb_response(wb, 'organization_order_export.xlsx')

    def template_excel(self, request):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = '宽表样例'
        ws1.append([
            '订单编号', '订单日期', '机构号', '线路号',
            '100元捆数', '50元捆数', '20元捆数', '10元捆数', '5元捆数',
            '1元包数', '0.5元包数', '0.1元包数', '订单状态', '备注',
        ])
        ws1.append(['ORD20260410001', '2026-04-10', 'ORG001', 'R001', 10, 2, 0, 0, 1, 0, 0, 0, 'NEW', '宽表样例'])

        ws2 = wb.create_sheet('长表样例')
        ws2.append(['订单编号', '订单日期', '机构号', '线路号', '面额', '数量', '金额', '订单状态', '备注'])
        ws2.append(['ORD20260410002', '2026-04-10', 'ORG001', 'R001', 100, 10, '', 'NEW', '数量导入样例'])
        ws2.append(['ORD20260410003', '2026-04-10', 'ORG001', 'R001', 1, '', 1500, 'NEW', '金额导入样例(按封装规格换算)'])

        return self._wb_response(wb, 'organization_order_template.xlsx')

    @staticmethod
    def _wb_response(wb, filename):
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename={filename}'
        return response


def _to_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value), '%Y-%m-%d').date()


def _normalize_denomination(value) -> str:
    d = Decimal(str(value)).normalize()
    txt = format(d, 'f').rstrip('0').rstrip('.')
    return txt if txt else '0'


def _get_configured_denominations() -> list[str]:
    configured = {
        _normalize_denomination(support.denomination)
        for support in StationDenominationSupport.objects.filter(enabled=True)
    }
    return [denom for denom in DENOMINATION_TO_FIELD if denom in configured]


def _validate_denomination_bound(denom_text: str):
    denomination = Decimal(denom_text)
    currency_type = CurrencyType.COIN if denomination < Decimal('5') else CurrencyType.BANKNOTE
    spec_exists = DenominationPackagingSpec.objects.filter(
        currency_type=currency_type,
        denomination=denomination,
        enabled=True,
    ).exists()
    support_exists = StationDenominationSupport.objects.filter(
        currency_type=currency_type,
        denomination=denomination,
        enabled=True,
    ).exists()
    if not spec_exists or not support_exists:
        missing_parts = []
        if not spec_exists:
            missing_parts.append('面额封装规格')
        if not support_exists:
            missing_parts.append('工位支持面额')
        raise ValueError(f'面额 {DENOMINATION_LABEL_MAP.get(denom_text, denom_text)} 未完成绑定：缺少{"、".join(missing_parts)}配置')


def _resolve_quantity_from_row(denom_text: str, qty_raw, amount_raw) -> int:
    if qty_raw not in (None, ''):
        return int(qty_raw)
    if amount_raw in (None, ''):
        raise ValueError(f'面额 {DENOMINATION_LABEL_MAP.get(denom_text, denom_text)} 未填写数量或金额')

    denomination = Decimal(denom_text)
    currency_type = CurrencyType.COIN if denomination < Decimal('5') else CurrencyType.BANKNOTE
    spec = DenominationPackagingSpec.objects.filter(
        currency_type=currency_type,
        denomination=denomination,
        enabled=True,
    ).first()
    if not spec:
        raise ValueError(f'面额 {DENOMINATION_LABEL_MAP.get(denom_text, denom_text)} 缺少启用的面额封装规格')

    units = int(spec.units_per_package)
    if units <= 0:
        raise ValueError(f'面额 {DENOMINATION_LABEL_MAP.get(denom_text, denom_text)} 封装规格单位必须大于0')

    amount = Decimal(str(amount_raw))
    qty = amount / (denomination * units)
    if qty != qty.to_integral_value():
        raise ValueError(
            f'面额 {DENOMINATION_LABEL_MAP.get(denom_text, denom_text)} 金额 {amount_raw} '
            f'不能按每包(捆){units}张(枚)整除换算为整数包(捆)数'
        )
    return int(qty)
